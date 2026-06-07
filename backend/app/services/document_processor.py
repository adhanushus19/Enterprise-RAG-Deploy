import os
import logging
from datetime import datetime
from typing import List, Dict, Any, Tuple, Optional
import pypdf
import docx
import pptx
import openpyxl
import pandas as pd
from pydantic import BaseModel, Field
from openai import OpenAI
from app.config import settings

logger = logging.getLogger(__name__)

class ExtractedMetadata(BaseModel):
    author: Optional[str] = Field(None, description="The author or creator of the document")
    department: Optional[str] = Field(None, description="The department or team associated with the document (e.g., HR, Finance, Engineering, Legal)")
    creation_date: Optional[str] = Field(None, description="The creation date of the document in YYYY-MM-DD format if mentioned, else null")
    tags: List[str] = Field(default_factory=list, description="A list of 3-5 keywords or tags summarizing the document topics")
    doc_type: str = Field(description="The document type (pdf, docx, pptx, xlsx)")


class DocumentProcessor:
    """
    Service responsible for parsing multiple enterprise document types (PDF, DOCX, PPTX, XLSX)
    and extracting content, page/cell locations, and semantic metadata.
    """

    def __init__(self):
        self.openai_client = None
        if settings.OPENAI_API_KEY:
            self.openai_client = OpenAI(api_key=settings.OPENAI_API_KEY)
        else:
            logger.warning("OPENAI_API_KEY is not set. Metadata extraction will fall back to rule-based defaults.")

    def parse_file(self, file_path: str, doc_type: str) -> List[Tuple[str, int]]:
        """
        Parses a file and returns list of (chunk_text, index) where index is page/slide/row/etc.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        doc_type = doc_type.lower()
        if doc_type == "pdf":
            return self._parse_pdf(file_path)
        elif doc_type == "docx":
            return self._parse_docx(file_path)
        elif doc_type == "pptx":
            return self._parse_pptx(file_path)
        elif doc_type in ["xlsx", "xls"]:
            return self._parse_xlsx(file_path)
        else:
            raise ValueError(f"Unsupported document type: {doc_type}")

    def _parse_pdf(self, file_path: str) -> List[Tuple[str, int]]:
        chunks = []
        try:
            reader = pypdf.PdfReader(file_path)
            for page_idx, page in enumerate(reader.pages):
                text = page.extract_text()
                if text and text.strip():
                    chunks.append((text.strip(), page_idx + 1))
        except Exception as e:
            logger.error(f"Failed to parse PDF {file_path}: {str(e)}", exc_info=True)
            raise e
        return chunks

    def _parse_docx(self, file_path: str) -> List[Tuple[str, int]]:
        chunks = []
        try:
            doc = docx.Document(file_path)
            current_paragraph_group = []
            chunk_idx = 1
            
            # DOCX does not have explicit pages, so we group paragraphs (approx 3-4 paragraphs per simulated page)
            for para in doc.paragraphs:
                if para.text.strip():
                    current_paragraph_group.append(para.text.strip())
                if len(current_paragraph_group) >= 4:
                    chunks.append(("\n\n".join(current_paragraph_group), chunk_idx))
                    current_paragraph_group = []
                    chunk_idx += 1
            
            if current_paragraph_group:
                chunks.append(("\n\n".join(current_paragraph_group), chunk_idx))
        except Exception as e:
            logger.error(f"Failed to parse DOCX {file_path}: {str(e)}", exc_info=True)
            raise e
        return chunks

    def _parse_pptx(self, file_path: str) -> List[Tuple[str, int]]:
        chunks = []
        try:
            prs = pptx.Presentation(file_path)
            for slide_idx, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                if slide_text:
                    chunks.append(("\n".join(slide_text), slide_idx + 1))
        except Exception as e:
            logger.error(f"Failed to parse PPTX {file_path}: {str(e)}", exc_info=True)
            raise e
        return chunks

    def _parse_xlsx(self, file_path: str) -> List[Tuple[str, int]]:
        chunks = []
        try:
            xls = pd.ExcelFile(file_path)
            for sheet_name in xls.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                if df.empty:
                    continue
                # Transform each row into a structured text chunk
                sheet_data = []
                for idx, row in df.iterrows():
                    row_repr = f"Sheet: {sheet_name} | Row {idx+1}: " + ", ".join([f"{col}={val}" for col, val in row.items() if pd.notna(val)])
                    sheet_data.append(row_repr)
                
                # Group rows to prevent tiny single-row vectors
                for i in range(0, len(sheet_data), 10):
                    group = sheet_data[i:i+10]
                    chunks.append(("\n".join(group), i // 10 + 1))
        except Exception as e:
            logger.error(f"Failed to parse XLSX {file_path}: {str(e)}", exc_info=True)
            raise e
        return chunks

    def extract_metadata(self, first_chunk_text: str, filename: str, doc_type: str) -> Dict[str, Any]:
        """
        Uses OpenAI LLM to analyze the initial content and extract structured metadata.
        Falls back to rule-based analysis if the API key is not present.
        """
        if self.openai_client:
            try:
                system_prompt = (
                    "You are an expert document analyzer. Extract the following metadata keys in JSON format:\n"
                    "- author: Creator of the document (name or role), if discernible.\n"
                    "- department: HR, Finance, Engineering, Legal, Operations, Sales, Marketing, or similar.\n"
                    "- creation_date: Date in YYYY-MM-DD format if mentioned, else return null.\n"
                    "- tags: 3 to 5 relevant tags summarizing the primary topics of the content.\n"
                    "- doc_type: The file format (must be one of: pdf, docx, pptx, xlsx).\n\n"
                    "Do not make up information. If a field is not present or cannot be inferred, return null."
                )
                
                prompt = f"Filename: {filename}\nFile Type: {doc_type}\n\nContent Sample:\n{first_chunk_text[:3000]}"
                
                # Use structured completion endpoint
                response = self.openai_client.beta.chat.completions.parse(
                    model="gpt-4o-mini",
                    messages=[
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": prompt}
                    ],
                    response_format=ExtractedMetadata,
                    temperature=0.0
                )
                
                parsed = response.choices[0].message.parsed
                creation_date_val = None
                if parsed.creation_date:
                    try:
                        creation_date_val = datetime.strptime(parsed.creation_date, "%Y-%m-%d")
                    except ValueError:
                        pass
                
                return {
                    "author": parsed.author or "Unknown",
                    "department": parsed.department or "General",
                    "creation_date": creation_date_val,
                    "tags": parsed.tags or [],
                    "doc_type": doc_type
                }
            except Exception as e:
                logger.error(f"LLM metadata extraction failed, falling back: {str(e)}")

        # Fallback implementation
        tags = ["ingested"]
        if doc_type:
            tags.append(doc_type.lower())
        
        return {
            "author": "System Ingest",
            "department": "General",
            "creation_date": datetime.now(),
            "tags": tags,
            "doc_type": doc_type
        }
